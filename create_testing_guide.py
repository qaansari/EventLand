import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, title="NOTE:", bg_hex="F0F4F8", border_hex="1E3A8A"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    # Title color determination
    if border_hex == "1E3A8A" or border_hex == "2563EB":
        title_color = RGBColor(30, 58, 138)
    elif border_hex == "DC2626":
        title_color = RGBColor(220, 38, 38)
    elif border_hex == "059669":
        title_color = RGBColor(5, 150, 105)
    else:
        title_color = RGBColor(180, 83, 9)

    run_title = p.add_run(f"{title} ")
    run_title.bold = True
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = title_color
    
    run_text = p.add_run(text)
    run_text.font.name = 'Calibri'
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = RGBColor(31, 41, 55)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_test_case(doc, tc_id, title, steps, expected_result, note=None, note_title="PRO TIP:"):
    h2 = doc.add_heading(level=2)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)
    r = h2.add_run(f"{tc_id}: {title}")
    r.font.color.rgb = RGBColor(15, 118, 110) # Teal
    r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    for i, step in enumerate(steps, 1):
        r_step_num = p.add_run(f"{i}. ")
        r_step_num.bold = True
        r_step_num.font.size = Pt(10.5)
        # Parse bold formatting if embedded in tuple or text
        if isinstance(step, tuple) or isinstance(step, list):
            for part, is_bold in step:
                r_part = p.add_run(part)
                r_part.bold = is_bold
                r_part.font.size = Pt(10.5)
        else:
            r_part = p.add_run(str(step))
            r_part.font.size = Pt(10.5)
        if i < len(steps):
            p.add_run("\n")

    p_exp = doc.add_paragraph()
    p_exp.paragraph_format.left_indent = Inches(0.25)
    p_exp.paragraph_format.space_before = Pt(2)
    p_exp.paragraph_format.space_after = Pt(6)
    r_exp_title = p_exp.add_run("Expected Result: ")
    r_exp_title.bold = True
    r_exp_title.font.color.rgb = RGBColor(30, 27, 75)
    r_exp_title.font.size = Pt(10.5)

    r_exp = p_exp.add_run(expected_result)
    r_exp.font.size = Pt(10.5)
    r_exp.font.color.rgb = RGBColor(51, 65, 85)

    if note:
        add_callout(doc, note, title=note_title, bg_hex="FFFBEB", border_hex="D97706")

def build_document():
    doc = docx.Document()
    
    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles setup
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

    # Palette
    PRIMARY = RGBColor(30, 27, 75)     # Deep Indigo (#1E1B4B)
    SECONDARY = RGBColor(15, 118, 110) # Teal (#0F766E)
    DARK_TEXT = RGBColor(15, 23, 42)   # Slate 900

    # -------------------------------------------------------------
    # DOCUMENT TITLE / HEADER BLOCK
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("EVENTLAND PLATFORM")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("End-to-End Testing Guide & Step-by-Step QA Manual (From Scratch)")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(16)
    run_sub.font.bold = True
    run_sub.font.color.rgb = SECONDARY

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(4)
    run_meta = meta_p.add_run("Platform Release: Phase 1 Final / Live Portal Edition  |  Target: .NET 10 Web API + React 19 (Vite)")
    run_meta.font.size = Pt(10)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    desc_p = doc.add_paragraph()
    desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc_p.paragraph_format.space_before = Pt(0)
    desc_p.paragraph_format.space_after = Pt(20)
    run_desc = desc_p.add_run("Comprehensive test execution manual covering all platform features across Super Admin, Admin, Organizer, Attendee/Customer roles, PayPro v2 Financial Switch & Gateway, Automated Webhooks, Background Reconciliation, Admin Bookings Hub, Cloudflare Turnstile Bot Defense, and SkiaSharp Pixel Sanitization.")
    run_desc.font.size = Pt(10.5)
    run_desc.font.italic = True

    # Horizontal Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(16)
    p_div_run = p_div.add_run("_________________________________________________________________________________")
    p_div_run.font.color.rgb = RGBColor(203, 213, 225)

    # -------------------------------------------------------------
    # SECTION 1: OVERVIEW & PREREQUISITES
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("1. Document Overview & Environment Setup")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("This manual provides a zero-to-hero, step-by-step procedure to test the entire ")
    r_bold = p.add_run("EventLand")
    r_bold.bold = True
    p.add_run(" ticketing and event management ecosystem from scratch. It verifies core business logic including user role authorization, interactive seat reservations, SignalR real-time locks, PayPro v2 financial switch operations, 1Pay instant online payment checkout, direct bank transfer manual verification, autonomous 30-minute hold expiration, automated webhook processing, background reconciliation, and digital E-Ticket QR code issuance.")

    add_callout(doc, 
                "Before starting test execution, ensure both backend API (.NET 10 Web API at localhost:5000 / https://localhost:7147) and frontend SPA (React 19 + Vite at localhost:5173) are running. PayPro sandbox credentials (Username, Password, ClientSecret, MerchantId) are configured in appsettings.Development.json under the 'PayPro' configuration section.", 
                title="PREREQUISITE CHECK:", 
                bg_hex="EFF6FF", 
                border_hex="2563EB")

    # Table of Environment details
    doc.add_heading("System Environment & Default Test Accounts", level=2)
    
    tbl_env = doc.add_table(rows=5, cols=3)
    tbl_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_env)

    headers = ["Role / Component", "URL / Access Details", "Default Credentials / Notes"]
    for i, h_text in enumerate(headers):
        cell = tbl_env.cell(0, i)
        set_cell_background(cell, "1E1B4B")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    env_data = [
        ("Super Admin", "http://localhost:5173 (Admin Portal)", "admin@eventland.pk / SuperAdmin123!"),
        ("Admin", "http://localhost:5173 (Admin Portal)", "admin.qa@eventland.pk (Created via SuperAdmin)"),
        ("Organizer", "http://localhost:5173 (Organizer Portal)", "organizer.qa@eventland.pk or Self-Register"),
        ("Attendee / Customer", "http://localhost:5173 (Public Web App)", "Guest or Self-Registered Customer Account")
    ]

    for row_idx, data in enumerate(env_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = tbl_env.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 2: SUPER ADMIN TESTING
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("2. Super Admin Role - System Initialization & Governance")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.add_run("The Super Admin holds supreme authority over the EventLand platform. Responsibilities include system security, global role assignments, active bank account management, bank downtime maintenance locks, locations (countries/cities), PayPro financial switch administration, FAQs, bot defense governance, and global audit.")

    # Test Case 1.1
    add_test_case(doc, "Test Case 1.1", "Super Admin Authentication & Initial Login", [
        [("Open the application in your browser at ", False), ("http://localhost:5173", True), (".", False)],
        [("Click ", False), ("Login / Register", True), (" in the top navigation header.", False)],
        [("Enter Email: ", False), ("admin@eventland.pk", True), (" and Password: ", False), ("SuperAdmin123!", True), (".", False)],
        [("Complete Cloudflare Turnstile verification challenge if prompted.", False)],
        [("Click ", False), ("Sign In", True), (".", False)]
    ], "JWT Bearer token is issued and stored in browser localStorage. Super Admin badge and 'Admin Portal' navigation link appear in header. Session is authenticated.")

    # Test Case 1.2
    add_test_case(doc, "Test Case 1.2", "System Active Bank Account & Maintenance Notice Setup", [
        [("Navigate to ", False), ("Admin Portal -> Bank Accounts", True), (" tab.", False)],
        [("Click ", False), ("Add New Bank Account", True), (" (or Edit existing).", False)],
        [("Verify Static Modal Constraint: Click outside the modal on the darkened background backdrop overlay. Verify that the modal does NOT close or dismiss (static modal behavior ensures in-progress financial inputs are not accidentally lost).", False)],
        [("Enter Bank Name: ", False), ("United Bank Limited (UBL)", True), (", Account Title: ", False), ("Event Land Official Pvt Ltd", True), (", Account Number: ", False), ("0123456789", True), (", IBAN: ", False), ("PK36UNIL0109000123456789", True), (", and Branch Code: ", False), ("0981", True), (".", False)],
        [("Upload Bank QR Code image file (e.g. UBL Raast QR Code).", False)],
        [("Toggle ", False), ("Is Maintenance Mode", True), (" to ON. Enter Maintenance Notice: ", False), ("'Bank network scheduled maintenance in progress.'", True)],
        [("Save the record. Verify public checkout modal displays advisory banner and locks ticket payment. Then toggle maintenance back to OFF.", False)]
    ], "Bank account modal enforces static backdrop constraint (clicks on overlay do not dismiss dialog; closes strictly via 'X' or 'Cancel'). Account details saved in DB table 'BankAccounts'. Active bank endpoint '/api/bank-accounts/active' returns details. Public checkout reflects live maintenance notice.")

    # Test Case 1.3
    add_test_case(doc, "Test Case 1.3", "User Management & Role Promotion (Admin & Organizer Creation)", [
        [("Navigate to ", False), ("Admin Portal -> Users & Roles", True), (" tab.", False)],
        [("Click ", False), ("Create User", True), (".", False)],
        [("Create Admin Account: Name: ", False), ("QA Operations Admin", True), (", Email: ", False), ("admin.qa@eventland.pk", True), (", Role: ", False), ("Admin", True), (", Password: ", False), ("Admin123!@#", True), (".", False)],
        [("Create Organizer Account: Name: ", False), ("Prime Events Pakistan", True), (", Email: ", False), ("organizer.qa@eventland.pk", True), (", Role: ", False), ("Organizer", True), (", Password: ", False), ("Organizer123!@#", True), (".", False)],
        [("Test Account Lockout Defense: Attempt 5 incorrect logins for an account. Verify account temporarily locks out for 15 minutes.", False)]
    ], "Users successfully registered in DB with hashed passwords. Password complexity and unique email/phone validations enforced. Temporary lockout triggers on 5 consecutive failures.")

    # Test Case 1.4
    add_test_case(doc, "Test Case 1.4", "Location (Countries & Cities) & Metadata Management", [
        [("Navigate to ", False), ("Admin Portal -> Countries", True), (" tab. Ensure Pakistan (Code: PK, Dial: +92) is active.", False)],
        [("Navigate to ", False), ("Cities", True), (" tab. Add or verify cities: ", False), ("Karachi, Lahore, Islamabad, Rawalpindi", True), (".", False)],
        [("Navigate to ", False), ("Tags / Categories", True), (" tab. Add/verify tags: ", False), ("Concerts, Theatre, Comedy, Festivals, Tech Conferences", True), (".", False)],
        [("Navigate to ", False), ("FAQs", True), (" tab. Add frequently asked question regarding ticket refund and seat holds.", False)]
    ], "Metadata records saved in DB. Homepage location dropdown and category pill filters reflect newly added cities and tags dynamically.")

    # Test Case 1.5 (NEW)
    add_test_case(doc, "Test Case 1.5", "Cloudflare Turnstile Bot Defense & Rate Limiting Verification", [
        [("Open the public landing page in an Incognito window.", False)],
        [("Locate the security verification widget (or open Login modal).", False)],
        [("Complete the Cloudflare Turnstile challenge checkbox.", False)],
        [("Observe that once passed, the CAPTCHA widget unmounts cleanly and is automatically hidden across the entire application for the active session (persisted via sessionStorage and window event bus).", False)],
        [("Verify backend verification endpoint: ", False), ("POST /api/captcha/verify", True), (" returns success token.", False)],
        [("Test Rate Limiting: Send 31 rapid requests to ", False), ("POST /api/auth/login", True), (" within 60 seconds from the same IP.", False)]
    ], "Captcha automatically unmounts across all components (AuthModal, Footer, App). Rate limiter enforces 30 requests/min per IP on login returning HTTP 429 Too Many Requests.",
    note="Turnstile auto-hide prevents repetitive user friction while per-IP rate limiting halts bot brute-force and credential stuffing.",
    note_title="SECURITY ARCHITECTURE:")

    # Test Case 1.6 (NEW)
    add_test_case(doc, "Test Case 1.6", "PayPro Gateway Switch Administration (PayProAdminPanel)", [
        [("Navigate to ", False), ("Admin Portal -> PayPro Gateway", True), (" tab (visible exclusively to SuperAdmin).", False)],
        [("Sub-Test 1.6.1 (Paid Orders Report - GPO): Click 'Paid Orders Report'. Set date range (e.g., last 30 days). Click 'Apply Filters'. Inspect paginated results, transaction references, amounts paid, and summary revenue metrics.", False)],
        [("Sub-Test 1.6.2 (Consumers Management): Click 'Consumers' sub-tab. Click 'Register New Consumer'. Enter Consumer ID (e.g. 'CUST-001'), Name, Mobile, Email, and Address. Save record. Click 'Batch CSV Import' and verify bulk consumer template parsing.", False)],
        [("Sub-Test 1.6.3 (Reconciliation & Ops): Click 'Reconciliation & Ops' sub-tab. Click 'Trigger Manual Sweep' (older than 15 minutes). Verify automated status polling across PayPro switch.", False)],
        [("Sub-Test 1.6.4 (Order Inspector & Status Override): Enter an Order Number (e.g. 'EVL-10023'). Click 'Inspect Live Status' to fetch live PayPro data via /v2/ppro/ggosboi. Test manual 'Mark Order As Paid' (moap) or 'Mark Order As Blocked' (moab) overrides.", False)]
    ], "PayProAdminPanel renders real-time switch data from EventLand.Modules.PayPro. Consumers registered, reports loaded, and reconciliation sweeps execute without exceptions.",
    note="PayPro switch operations allow Super Admins to audit live gateway settlements, register corporate consumers, and resolve disputed transactions directly.",
    note_title="FINANCIAL SWITCH FEATURE:")

    # -------------------------------------------------------------
    # SECTION 3: ADMIN TESTING
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("3. Admin Role - Platform Operations, Multi-Show Scheduling & Booking Oversight")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.add_run("Admins manage platform venues, auditorium layout charts, seating zones, artist profiles, multi-show schedules, event publishing, order processing, and bank transfer manual payment verification.")

    # Test Case 2.1
    add_test_case(doc, "Test Case 2.1", "Venue & Interactive Auditorium Layout Configuration", [
        [("Log in as Admin (", False), ("admin.qa@eventland.pk", True), ("). Navigate to ", False), ("Admin Portal -> Venues & Layouts", True), (".", False)],
        [("Click ", False), ("Add Venue", True), (". Enter Venue Name: ", False), ("Arts Council Karachi", True), (", City: ", False), ("Karachi", True), (", Address: ", False), ("M.R. Kiyani Road, Saddar, Karachi", True), (".", False)],
        [("Navigate to ", False), ("Auditorium Charts", True), (" tab. Click ", False), ("Add Auditorium Layout", True), (".", False)],
        [("Enter Auditorium Name: ", False), ("Main Auditorium", True), (", Rows: ", False), ("10", True), (", Columns: ", False), ("20 (Capacity: 200 seats)", True), (".", False)],
        [("Define Seating Zones: Zone 1: ", False), ("VIP (Rows A-C, 20 seats/row)", True), (", Zone 2: ", False), ("Executive (Rows D-J, 20 seats/row)", True), (".", False)],
        [("Click ", False), ("Export Seating Chart PDF", True), (" to verify client-side PDF layout generation.", False)]
    ], "Venue and auditorium layout saved. Seat grid generated in database with unique row and column IDs. Seating chart PDF exports cleanly.")

    # Test Case 2.2
    add_test_case(doc, "Test Case 2.2", "Artist Profile Management", [
        [("Navigate to ", False), ("Admin Portal -> Artists", True), (" tab.", False)],
        [("Click ", False), ("Add New Artist", True), (".", False)],
        [("Enter Artist Name: ", False), ("Atif Aslam", True), (", Genre: ", False), ("Pop / Rock", True), (", City: ", False), ("Lahore", True), (", Starting Rate: ", False), ("PKR 5,000,000", True), (", Bio: ", False), ("Internationally acclaimed vocalist and live performer.", True), (".", False)],
        [("Upload Artist Profile Photo (PNG/JPG). Verify image upload re-encoding.", False)],
        [("Click ", False), ("Save Artist", True), (". Verify artist appears in artist directory.", False)]
    ], "Artist record created with sanitized image asset. Available for linking to upcoming concert events.")

    # Test Case 2.3
    add_test_case(doc, "Test Case 2.3", "Global Event Publishing & SEO Slug Generation", [
        [("Navigate to ", False), ("Admin Portal -> Events", True), (" tab. Click ", False), ("Add New Event", True), (".", False)],
        [("Enter Title: ", False), ("Atif Aslam Live in Concert 2026", True), (", Category: ", False), ("Concerts", True), (", Venue: ", False), ("Arts Council Karachi", True), (".", False)],
        [("Set Start Date/Time and End Date/Time.", False)],
        [("Upload Event Banner image and Poster image.", False)],
        [("Toggle ", False), ("Is Published", True), (" to ON. Save the event.", False)],
        [("Verify SEO URL Slug: Check that the event receives a clean slug URL (e.g. ", False), ("/event/atif-aslam-live-in-concert-2026-1001", True), (").", False)]
    ], "Event saved in DB with status 'Live' / 'Published'. SEO slug generated automatically. Event is visible on homepage.")

    # Test Case 2.4 (NEW)
    add_test_case(doc, "Test Case 2.4", "Multi-Show Scheduling & Ticket Tier Row Range Mapping", [
        [("Edit the newly created event. Navigate to the ", False), ("Shows & Schedules", True), (" section.", False)],
        [("Add Show 1: ", False), ("Show Date/Time: Saturday 8:00 PM, Show Title: 'Saturday Night Live'", True), (".", False)],
        [("Add Show 2: ", False), ("Show Date/Time: Sunday 7:00 PM, Show Title: 'Sunday Grand Finale'", True), (".", False)],
        [("Navigate to ", False), ("Ticket Tiers", True), (" section.", False)],
        [("Create Tier 1: Name: ", False), ("VIP Passes", True), (", Price: ", False), ("PKR 6,000", True), (", Capacity: ", False), ("60", True), (", Row Range: ", False), ("A-C", True), (", Linked Show: ", False), ("Saturday Night Live", True), (".", False)],
        [("Create Tier 2: Name: ", False), ("Executive Passes", True), (", Price: ", False), ("PKR 3,000", True), (", Capacity: ", False), ("140", True), (", Row Range: ", False), ("D-J", True), (", Linked Show: ", False), ("Saturday Night Live", True), (".", False)],
        [("Save tiers. Verify composite database index on TicketTiers(EventId, EventShowId) ensures high-speed querying.", False)]
    ], "Multiple shows created via AdminEventShowsController. Ticket tiers linked to specific shows with row ranges. Public event detail page renders show selector dropdown.",
    note="Multi-show support allows a single event to run across multiple dates while keeping seat reservations segregated per show.",
    note_title="MULTI-SHOW SCHEDULING:")

    # Test Case 2.5 (NEW)
    add_test_case(doc, "Test Case 2.5", "Admin Bookings & E-Tickets Management Hub (AdminBookingsTab)", [
        [("Navigate to ", False), ("Admin Portal -> Bookings & Tickets", True), (" tab.", False)],
        [("Inspect Live KPI Metric Cards at top: ", False), ("Total Orders Placed, Total Tickets Sold, Total Confirmed Revenue", True), (" (PKR).", False)],
        [("Test Multi-Parameter Filters: Filter by Event (e.g. 'Atif Aslam Live'), Filter by Show, Filter by Payment Status (Paid, Pending, Confirmed, Cancelled).", False)],
        [("Test Live Keyword Search: Search by Booking Reference (e.g. 'EVL-'), Customer Name, or Customer Email.", False)],
        [("Test Data Export: Click ", False), ("'Export CSV'", True), (" and verify CSV download. Click ", False), ("'Export Excel (.xlsx)'", True), (" and verify formatted spreadsheet export.", False)],
        [("Inspect Payment Proof: Click the Eye icon on a pending bank transfer booking. Verify high-resolution payment slip preview modal.", False)],
        [("Click ", False), ("'Confirm Bank Payment'", True), (" on a pending order. Verify booking immediately updates to Confirmed/Paid and generates digital E-Ticket.", False)],
        [("Click ", False), ("'View Digital Ticket'", True), (" to open DigitalTicketModal. Click 'Download PDF Ticket' to verify client-side PDF export.", False)]
    ], "AdminBookingsTab provides end-to-end booking governance, KPI reporting, multi-filter search, CSV/Excel exports, and instant bank proof verification.",
    note="AdminBookingsTab isolates booking administration with dedicated performance optimizations and real-time revenue calculations.",
    note_title="ADMIN HUB ENHANCEMENT:")

    # -------------------------------------------------------------
    # SECTION 4: ORGANIZER TESTING
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("4. Organizer Role - Event Creation & Sales Analytics")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.add_run("Organizers manage self-service event registration, multi-show schedules, ticket tier allocations, dedicated revenue analytics, and attendee lists for their own events.")

    # Test Case 3.1
    add_test_case(doc, "Test Case 3.1", "Self-Service Event Creation Wizard", [
        [("Log in as Organizer (", False), ("organizer.qa@eventland.pk", True), (").", False)],
        [("Click ", False), ("List Your Event", True), (" in navbar or open ", False), ("Organizer Dashboard -> Create Event", True), (".", False)],
        [("Step 1 (Basic Details): Enter Title, Description, City, Category Tag.", False)],
        [("Step 2 (Schedule): Define Start/End Dates and configure multiple Show Schedules.", False)],
        [("Step 3 (Ticket Tiers & Seating): Add VIP and Standard tiers with prices and quantities.", False)],
        [("Step 4 (Media Upload): Upload Banner, Poster, and promotional photos.", False)],
        [("Submit Event. Verify event is created under the organizer's account ID.", False)]
    ], "Event registered successfully under Organizer ownership. Listed exclusively in Organizer Dashboard.")

    # Test Case 3.2
    add_test_case(doc, "Test Case 3.2", "BOLA / IDOR Authorization Security Defenses", [
        [("Logged in as Organizer A, capture an Event ID owned by Admin or Organizer B.", False)],
        [("Attempt to invoke edit or delete API endpoints (e.g. ", False), ("PUT /api/admin/events/{foreign_id}", True), (" or ", False), ("DELETE /api/admin/events/{foreign_id}", True), (") via API client or modified frontend state.", False)]
    ], "Backend rejects the request with HTTP 403 Forbidden ('Unauthorized access to event'). Organizers cannot access or modify events owned by other accounts.",
    note="Single-query SQL-level authorization scoping appends organizerId check to all DB query filters, preventing Broken Object Level Authorization (BOLA).",
    note_title="SECURITY HARDENING RULE:")

    # Test Case 3.3
    add_test_case(doc, "Test Case 3.3", "Organizer Sales & Booking Analytics Dashboard", [
        [("Open ", False), ("Organizer Dashboard", True), (".", False)],
        [("Review Overview Analytics: Total Revenue, Tickets Sold, Total Page Views, Sales by Tier.", False)],
        [("Click ", False), ("Artist Bookings", True), (" tab to review linked artist performance dates.", False)],
        [("Click ", False), ("Export Attendee List", True), (" to download attendee contact and seat booking records.", False)]
    ], "Analytics accurately reflect DB totals for the organizer's events only. Attendee export downloads cleanly.")

    # -------------------------------------------------------------
    # SECTION 5: ATTENDEE TESTING
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("5. Attendee Role - Discovery, Hold, PayPro Gateway & Bank Transfer")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.add_run("Attendees discover events via database-wide search and city filters, select seats with real-time SignalR locks, place 30-minute holds, checkout via PayPro 1Pay online gateway or direct bank transfer, track order status in real-time, and access digital QR E-Tickets.")

    # Test Case 4.1
    add_test_case(doc, "Test Case 4.1", "Homepage Discovery, Database-Wide Search & Category Filters", [
        [("Access public homepage at ", False), ("http://localhost:5173", True), (".", False)],
        [("Filter events by City (e.g. ", False), ("Karachi", True), (") using the top city dropdown.", False)],
        [("Filter by Category (e.g. ", False), ("Concerts", True), (") using category pill buttons.", False)],
        [("Type search query into the search bar (e.g. ", False), ("'Atif'", True), ("). Verify debounced search queries the entire database via ", False), ("GET /api/events?search=Atif", True), (".", False)],
        [("Click an event card to navigate to its SEO slug detail page.", False)]
    ], "Filtered events update smoothly. Database-wide search retrieves matching records across all pages. Slug URL renders cleanly.")

    # Test Case 4.2
    add_test_case(doc, "Test Case 4.2", "Interactive Seat Selection & Real-Time SignalR Locks", [
        [("Open Event Detail page and click ", False), ("Select Seats", True), (".", False)],
        [("In Browser 1 (Normal Window), select Show Date and click Seat ", False), ("A-5 (VIP)", True), (".", False)],
        [("Simultaneously open the same event seating map in Browser 2 (Incognito Window).", False)],
        [("Observe Seat A-5 status in Browser 2 in real-time.", False)]
    ], "SignalR seating hub ('/hubs/seating') broadcasts seat lock instantly. Browser 2 displays Seat A-5 as 'Reserved/Held' within milliseconds without page refresh.")

    # Test Case 4.3
    add_test_case(doc, "Test Case 4.3", "30-Minute Hold & Booking Reference Generation", [
        [("In Browser 1, enter customer details (Name, Email, Mobile).", False)],
        [("Click ", False), ("Proceed to Checkout", True), (".", False)],
        [("Observe generated unique Booking Reference code (e.g. ", False), ("EVL-894215", True), (").", False)],
        [("Verify the 30-minute hold countdown timer (PaymentExpiresAt) is actively ticking down in CheckoutModal.", False)]
    ], "Booking record created in DB with status 'Pending'. Seats locked for 30 minutes. Ephemeral Redis lock transitioned to DB hold.")

    # Test Case 4.4
    add_test_case(doc, "Test Case 4.4", "PayPro 1Pay Instant Online Gateway Checkout", [
        [("In Step 2 of Checkout Modal, select ", False), ("PayPro Online Gateway ⚡", True), (".", False)],
        [("Choose payment channel (e.g. ", False), ("EasyPaisa / JazzCash / Cards", True), (" or ", False), ("PayPro Instant QR Code", True), (").", False)],
        [("Click ", False), ("Proceed to PayPro Online Gateway →", True), (".", False)],
        [("System calls ", False), ("POST /api/payments/paypro/checkout", True), (" which authenticates against PayPro sandbox and creates order via /v2/ppro/co.", False)],
        [("Observe generated PayPro Consumer Voucher / OTC Number and Click2Pay portal link.", False)]
    ], "PayPro order generated with unique OTC Voucher Number and hosted checkout URL. Step 3 renders PayProStatusTracker.")

    # Test Case 4.5
    add_test_case(doc, "Test Case 4.5", "Direct Bank Transfer Checkout & Payment Proof Submission", [
        [("In Step 2 of Checkout Modal, select ", False), ("Direct Bank Transfer 🏛️", True), (".", False)],
        [("Review verified receiving bank details (UBL Bank, Account Title, IBAN, UBL QR Code).", False)],
        [("Path A (Receipt Image Upload): Upload payment receipt screenshot (JPEG/PNG). Verify SkiaSharp sanitizes raster pixels.", False)],
        [("Path B (Fallback Bank Details): Enter Sender Account Title, Sender Bank Name, and Sender Account Last 4 Digits.", False)],
        [("Click ", False), ("Submit Payment Proof", True), (".", False)]
    ], "Payment proof recorded. Booking status updates to 'PendingVerification'. Confirmation modal displayed with booking reference.")

    # Test Case 4.6
    add_test_case(doc, "Test Case 4.6", "Autonomous 30-Minute Hold Expiration (Negative Test)", [
        [("Place a seat hold for a booking and note the booking reference.", False)],
        [("Do NOT submit payment proof or complete gateway payment.", False)],
        [("Allow 30 minutes to elapse (or observe background worker execution).", False)],
        [("Verify PendingBookingExpiryService background worker (runs every 60s) cancels the booking.", False)]
    ], "Booking marked as 'Expired' / 'Cancelled'. SoldCount decremented. Held seats immediately returned to 'Available' pool.",
    note="PendingBookingExpiryService runs on composite index IX_Bookings_PaymentStatus_PaymentExpiresAt for sub-second cleanup.",
    note_title="BACKGROUND EXPIRY WORKER:")

    # Test Case 4.7 (NEW)
    add_test_case(doc, "Test Case 4.7", "Interactive PayPro Status Tracker Widget (PayProStatusTracker)", [
        [("On Step 3 of PayPro Checkout, observe the ", False), ("PayProStatusTracker", True), (" widget.", False)],
        [("Verify live pulsing badge indicating automated polling every 5 seconds.", False)],
        [("Click the manual refresh button to trigger an immediate status query.", False)],
        [("Click 'Copy Order Number' button to verify 1-click clipboard copy.", False)],
        [("Click 'Pay Online via PayPro 1Pay Portal' to test hosted portal navigation.", False)]
    ], "PayProStatusTracker polls /api/paypro/order-status/{orderNumber} seamlessly. Stops auto-polling upon reaching terminal status (Paid/Blocked).")

    # Test Case 4.8 (NEW)
    add_test_case(doc, "Test Case 4.8", "PayPro Hosted Return Page (PayProReturnPage)", [
        [("Simulate user redirection back from PayPro hosted portal by opening URL: ", False), ("http://localhost:5173/payments/return?ordId=EVL-894215&status=paid&msg=Success", True), (".", False)],
        [("Observe PayProReturnPage automatically verifies payment status against backend APIs.", False)],
        [("Verify receipt details: Order Number, Amount Paid, Payment Status, and E-Ticket confirmation.", False)],
        [("Click 'View My Digital Tickets' to navigate directly to Attendee Dashboard.", False)]
    ], "PayProReturnPage verifies order state, presents user-friendly receipt, and provides 1-click navigation to digital tickets.")

    # Test Case 4.9 (NEW)
    add_test_case(doc, "Test Case 4.9", "Client-Side Bot Defense & Auto-Hide Lifecycle", [
        [("Open the application in a fresh browser session (clear sessionStorage).", False)],
        [("Open Login / Register modal and complete the Turnstile challenge.", False)],
        [("Close modal and scroll to the website footer. Verify newsletter captcha is already marked verified and hidden.", False)],
        [("Proceed to checkout an event ticket. Verify no disruptive captcha challenges block the checkout flow.", False)]
    ], "Cloudflare Turnstile auto-hide coordinates state across components via sessionStorage and custom events, ensuring zero checkout friction.")

    # -------------------------------------------------------------
    # SECTION 6: PAYMENT VERIFICATION & E-TICKET LIFECYCLE
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("6. Payment Verification, Automated Webhooks & E-Ticket Lifecycle")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.add_run("This section covers both automated online gateway confirmation (via inbound webhooks and background reconciliation) and manual bank transfer verification, followed by digital QR E-Ticket generation.")

    # Test Case 5.1
    add_test_case(doc, "Test Case 5.1", "Admin Review of Unpaid Payment Invoices & Payment Proofs", [
        [("Log in as Admin / SuperAdmin and navigate to ", False), ("Admin Portal -> Bookings & Tickets", True), (".", False)],
        [("Filter bookings by status: ", False), ("'Pending'", True), (" or search for booking reference (e.g. ", False), ("EVL-894215", True), (").", False)],
        [("Click Eye icon to inspect uploaded bank payment slip image in full-screen modal preview, or verify sender bank name and account digits.", False)]
    ], "Payment proof details display accurately with high-resolution image preview and sender fallback details.")

    # Test Case 5.2
    add_test_case(doc, "Test Case 5.2", "Manual Payment Approval & Digital QR E-Ticket Issuance", [
        [("Click ", False), ("'Confirm Bank Payment'", True), (" on the pending booking.", False)],
        [("Verify booking status updates to ", False), ("'Paid' / 'Confirmed'", True), (".", False)],
        [("Verify seats permanently update from 'Reserved' to 'Booked'.", False)],
        [("Verify background confirmation email dispatch and WhatsApp attendee share link generation.", False)]
    ], "Booking confirmed. Seats marked 'Booked'. Digital E-Ticket pass generated with unique scannable QR code.")

    # Test Case 5.3
    add_test_case(doc, "Test Case 5.3", "Attendee Digital E-Ticket Viewing & Gate Validation", [
        [("Log in as Attendee and navigate to ", False), ("Attendee Dashboard -> My Bookings", True), (".", False)],
        [("Click ", False), ("'View Digital E-Ticket'", True), (" on the confirmed booking.", False)],
        [("Inspect DigitalTicketModal for crisp QR code, booking reference, seat numbers, show time, and venue address.", False)],
        [("Click ", False), ("'Download PDF Ticket'", True), (" to verify high-resolution PDF ticket pass download.", False)]
    ], "Digital E-Ticket renders formatted pass with crisp QR Code and venue metadata. PDF ticket downloads cleanly.")

    # Test Case 5.4 (NEW)
    add_test_case(doc, "Test Case 5.4", "Automated PayPro Webhook Inbound Ingestion (POST /paypro/uis)", [
        [("Simulate PayPro IPN payment webhook callback by issuing HTTP POST request to ", False), ("http://localhost:5000/paypro/uis", True), (" (or /api/payments/paypro-ipn).", False)],
        [("Headers: ", False), ("Content-Type: application/json", True), (", Body: ", False), ('{"username":"EVENTLAND_USER","password":"EVENTLAND_PASSWORD","csvinvoiceids":"EVL-894215"}', True), (".", False)],
        [("Verify backend validates merchant credentials using constant-time comparison (CryptographicOperations.FixedTimeEquals).", False)],
        [("Verify response: HTTP 200 OK with JSON array ", False), ('[{"StatusCode":"00","InvoiceID":"EVL-894215","Description":"Invoice successfully marked as paid"}]', True), (".", False)],
        [("Test Idempotency: Replay the identical webhook payload. Verify response is still StatusCode '00' without duplicate tickets or errors.", False)],
        [("Verify database: Booking EVL-894215 automatically updates to 'Paid', seats update to 'Booked', and QR tickets are generated without manual admin action.", False)]
    ], "Inbound PayPro callback processes securely in constant time, records audit log in PayProCallbackLogs, confirms booking idempotently, and issues digital tickets.",
    note="Webhook processing is 100% idempotent: duplicate callbacks return success without duplicating tickets or records.",
    note_title="FINANCIAL SWITCH WEBHOOK:")

    # Test Case 5.5 (NEW)
    add_test_case(doc, "Test Case 5.5", "Automated PayPro Background Reconciliation Service", [
        [("Create an order via PayPro but simulate a dropped webhook callback.", False)],
        [("Wait for the scheduled execution of ", False), ("PayProReconciliationJob", True), (" background worker (runs every 10 minutes).", False)],
        [("Verify worker queries Orders table using composite index IX_Orders_Status_UpdatedAtUtc for pending orders older than threshold.", False)],
        [("Worker queries PayPro live status endpoint (/v2/ppro/ggosboi) for the pending order.", False)],
        [("Upon confirming paid status from PayPro, worker triggers PayProBookingPaidHandler to automatically confirm seats and issue tickets.", False)]
    ], "Automated reconciliation detects unconfirmed paid orders and reconciles them with PayPro switch without human intervention.")

    # -------------------------------------------------------------
    # SECTION 7: SECURITY HARDENING & THREAT PREVENTION
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("7. Security Hardening & Threat Prevention Defenses")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.add_run("EventLand implements enterprise-grade cybersecurity controls to protect user data, prevent financial tampering, and defend against OWASP Top 10 vulnerabilities.")

    # Test Case 6.1 (NEW)
    add_test_case(doc, "Test Case 6.1", "Malicious File Upload & Pixel Re-Encoding Defense (UploadController)", [
        [("Test Executable Header Defense: Rename an executable file (PE .exe with 'MZ' signature or Linux ELF with '\\x7fELF') to 'receipt.png'. Attempt upload via ", False), ("POST /api/upload", True), (".", False)],
        [("Verify response: Backend rejects upload with HTTP 400 Bad Request ('Malicious file signature detected').", False)],
        [("Test Script Steganography Defense: Create an image file containing embedded PHP/JS script tags ('<?php system($_GET[\"cmd\"]); ?>'). Attempt upload.", False)],
        [("Verify response: SkiaSharp raster pixel re-encoding decodes pure pixels and strips all polyglots, steganography, and EXIF metadata.", False)],
        [("Test Image Decompression Bomb Defense: Attempt upload of an image exceeding 4096 × 4096 px dimensions. Verify backend rejects with 400 Bad Request.", False)]
    ], "Zero raw stream copying: All uploads decoded into memory bitmaps and re-encoded to pure raster pixels. Executable signatures, polyglots, and decompression bombs rejected.",
    note="UploadController enforces strict folder allowlists (events, slips, qr_codes, organizers, users) and cryptographically random GUID-based filenames.",
    note_title="UPLOAD HARDENING:")

    # Test Case 6.2 (NEW)
    add_test_case(doc, "Test Case 6.2", "Per-IP Partitioned Rate Limiting Verification", [
        [("Test Login Policy: Execute 31 login attempts within 60s from the same IP address.", False)],
        [("Verify response: 31st request returns HTTP 429 Too Many Requests.", False)],
        [("Test Booking Policy: Execute 31 booking creation attempts within 60s from the same IP.", False)],
        [("Verify response: Exceeding 30 req/min triggers HTTP 429 Too Many Requests.", False)],
        [("Test Upload Policy: Execute 21 file uploads within 60s from the same IP. Verify 21st request returns HTTP 429.", False)]
    ], "Rate limiting partitions requests per client IP (respecting CF-Connecting-IP and X-Forwarded-For). Halts brute force, ticket scalping, and upload flooding.")

    # Test Case 6.3 (NEW)
    add_test_case(doc, "Test Case 6.3", "OWASP Top 10 Defenses & Header Protections", [
        [("Inspect HTTP Response Headers using browser DevTools Network tab or curl: Verify ", False), ("X-Frame-Options: DENY", True), (", ", False), ("X-Content-Type-Options: nosniff", True), (", and Content Security Policy with ", False), ("frame-ancestors 'none'", True), (".", False)],
        [("Verify iframe embedding is blocked to prevent clickjacking attacks.", False)],
        [("Test SQL Injection: Submit SQL injection payloads in search queries (e.g. \"' OR 1=1 --\"). Verify queries execute via parameterized EF Core queries with zero injection vulnerability.", False)],
        [("Verify Ngrok Interstitial Bypass: Verify frontend requests include 'ngrok-skip-browser-warning': 'true' header to ensure clean CORS support across tunneling environments.", False)]
    ], "HTTP security headers enforced. Clickjacking prevented. SQL injection neutralized by EF Core parameterized queries. Reverse-proxy headers configured cleanly.")

    # -------------------------------------------------------------
    # SECTION 8: FULL QA TEST MATRIX SUMMARY
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("8. Complete QA Verification Matrix Checklist")
    r1.font.color.rgb = PRIMARY
    r1.font.bold = True

    p = doc.add_paragraph()
    p.add_run("Use this structured execution checklist to log and verify QA test runs across all roles and system components:")

    tbl_matrix = doc.add_table(rows=32, cols=5)
    tbl_matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_matrix)

    m_headers = ["ID", "Target Role", "Feature Under Test", "Key Verification Criteria", "Status"]
    for i, h_text in enumerate(m_headers):
        cell = tbl_matrix.cell(0, i)
        set_cell_background(cell, "1E1B4B")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    matrix_rows = [
        ("TC-1.1", "Super Admin", "Authentication & JWT", "Login succeeds, Admin Portal link & SuperAdmin badge rendered", "[  ] Pass"),
        ("TC-1.2", "Super Admin", "Bank Account Setup", "Static modal prevents backdrop close; active bank saved & maintenance locks checkout", "[  ] Pass"),
        ("TC-1.3", "Super Admin", "User & Role Management", "Promote roles, password rules & 15-min lockout verified", "[  ] Pass"),
        ("TC-1.4", "Super Admin", "Locations & Metadata", "Countries, Cities, FAQs & Tags dynamically managed", "[  ] Pass"),
        ("TC-1.5", "Super Admin", "Turnstile & Rate Limits", "Captcha auto-hides across app; per-IP rate limits enforce 429", "[  ] Pass"),
        ("TC-1.6", "Super Admin", "PayPro Switch Admin", "GPO reports, Consumer CSV import, manual sweep & status overrides", "[  ] Pass"),
        ("TC-2.1", "Admin", "Venue & Layout Config", "Auditorium & Seating zones created; PDF seating chart exports", "[  ] Pass"),
        ("TC-2.2", "Admin", "Artist Management", "Artist bio and photo uploaded, sanitized and saved", "[  ] Pass"),
        ("TC-2.3", "Admin", "Event Publishing", "Event published with clean SEO URL slug", "[  ] Pass"),
        ("TC-2.4", "Admin", "Multi-Show Scheduling", "Multiple shows scheduled with tier row range mapping", "[  ] Pass"),
        ("TC-2.5", "Admin", "Admin Bookings Hub", "KPI metrics, multi-filters, CSV/Excel export & slip preview", "[  ] Pass"),
        ("TC-3.1", "Organizer", "Event Creation Wizard", "Multi-step wizard completes event and show registration", "[  ] Pass"),
        ("TC-3.2", "Organizer", "BOLA / IDOR Defense", "Foreign event edits blocked with 403 Forbidden", "[  ] Pass"),
        ("TC-3.3", "Organizer", "Sales Analytics", "Total sales & revenue match DB counts; attendee list exports", "[  ] Pass"),
        ("TC-4.1", "Attendee", "Discovery & Search", "Database-wide search & category filters update list smoothly", "[  ] Pass"),
        ("TC-4.2", "Attendee", "SignalR Live Seat Lock", "Seat locked in real-time across multiple browsers", "[  ] Pass"),
        ("TC-4.3", "Attendee", "30-Min Seat Hold", "EVL reference generated & countdown active in modal", "[  ] Pass"),
        ("TC-4.4", "Attendee", "PayPro Online Gateway", "PayPro 1Pay voucher generated & 1Pay portal link active", "[  ] Pass"),
        ("TC-4.5", "Attendee", "Bank Transfer Checkout", "Proof image or bank fallback details submitted cleanly", "[  ] Pass"),
        ("TC-4.6", "Attendee", "Hold Expiration", "Expired holds auto-cancel and return seats to available pool", "[  ] Pass"),
        ("TC-4.7", "Attendee", "PayPro Status Tracker", "Live polling widget tracks status & provides Click2Pay link", "[  ] Pass"),
        ("TC-4.8", "Attendee", "PayPro Return Page", "Hosted return page verifies order & renders digital receipt", "[  ] Pass"),
        ("TC-4.9", "Attendee", "Turnstile Bot Defense", "Turnstile verification passes seamlessly without checkout friction", "[  ] Pass"),
        ("TC-5.1", "Admin / Attendee", "Bank Proof Review", "Full-resolution receipt preview modal renders uploaded slip", "[  ] Pass"),
        ("TC-5.2", "Admin / Attendee", "Manual Ticket Issuance", "Payment approval marks seats booked & issues digital QR pass", "[  ] Pass"),
        ("TC-5.3", "Attendee", "Digital E-Ticket View", "Attendee views scannable QR ticket & downloads PDF ticket", "[  ] Pass"),
        ("TC-5.4", "Gateway Switch", "PayPro IPN Webhook", "Inbound POST /paypro/uis confirms orders idempotently in real-time", "[  ] Pass"),
        ("TC-5.5", "Gateway Switch", "Auto-Reconciliation", "10-min background worker resolves unpaid orders with PayPro API", "[  ] Pass"),
        ("TC-6.1", "Security", "Malware & Upload Defense", "Executable PE/ELF headers rejected; SkiaSharp re-encodes pure pixels", "[  ] Pass"),
        ("TC-6.2", "Security", "IP Rate Limiting", "Per-IP rate limiting enforces 429 on login/booking/upload floods", "[  ] Pass"),
        ("TC-6.3", "Security", "OWASP & Security Headers", "Clickjacking blocked; EF Core SQL injection immune; CORS hardened", "[  ] Pass")
    ]

    for row_idx, data in enumerate(matrix_rows, start=1):
        if row_idx >= len(tbl_matrix.rows):
            tbl_matrix.add_row()
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = tbl_matrix.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = PRIMARY
            elif col_idx == 4:
                r.bold = True
                r.font.color.rgb = SECONDARY

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Save logic: save to current workspace directory and e:\EventLand if accessible
    script_dir = os.path.dirname(os.path.abspath(__file__))
    primary_output_path = os.path.join(script_dir, "EventLand_Full_Role_Testing_Guide.docx")
    try:
        doc.save(primary_output_path)
        print(f"Successfully generated Word document at: {primary_output_path}")
    except PermissionError:
        fallback_path = os.path.join(script_dir, "EventLand_Full_Role_Testing_Guide_Updated.docx")
        doc.save(fallback_path)
        print(f"Notice: 'EventLand_Full_Role_Testing_Guide.docx' is currently open in Microsoft Word and locked.")
        print(f"Saved latest updated document as: {fallback_path}")

    alt_dir = r"e:\EventLand"
    if os.path.exists(alt_dir) and os.path.isdir(alt_dir):
        alt_output_path = os.path.join(alt_dir, "EventLand_Full_Role_Testing_Guide.docx")
        try:
            doc.save(alt_output_path)
            print(f"Also saved copy at: {alt_output_path}")
        except Exception as e:
            print(f"Note: Could not write copy to {alt_output_path}: {e}")

if __name__ == "__main__":
    build_document()
